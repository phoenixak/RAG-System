"""Turn uploaded files into documents made of titled sections (PDF, DOCX, Markdown, TXT, CSV)."""

import csv
import hashlib
import io
import re
from dataclasses import dataclass, field
from pathlib import Path

SUPPORTED = {".pdf", ".docx", ".md", ".markdown", ".txt", ".csv"}
MAX_BYTES = 20 * 1024 * 1024


class UnsupportedFile(ValueError):
    pass


@dataclass
class Section:
    text: str
    heading: str | None = None
    page: int | None = None


@dataclass
class Document:
    doc_id: str
    source: str  # file name shown in citations
    title: str
    sections: list[Section] = field(default_factory=list)

    @property
    def text(self) -> str:
        return "\n\n".join(s.text for s in self.sections)


def doc_id_for(content: bytes) -> str:
    """Same bytes → same id, which is how re-uploads are detected."""
    return hashlib.sha256(content).hexdigest()[:16]


def load_path(path: str | Path) -> Document:
    p = Path(path)
    return load_bytes(p.read_bytes(), p.name)


def load_bytes(content: bytes, filename: str) -> Document:
    ext = Path(filename).suffix.lower()
    if ext not in SUPPORTED:
        raise UnsupportedFile(
            f"{ext or 'files without an extension'} not supported; use {', '.join(sorted(SUPPORTED))}"
        )
    if len(content) > MAX_BYTES:
        raise UnsupportedFile(f"{filename} is larger than {MAX_BYTES // (1024 * 1024)} MB")
    if ext == ".pdf":
        sections = _pdf(content)
    elif ext == ".docx":
        sections = _docx(content)
    elif ext == ".csv":
        sections = _csv(content)
    else:
        sections = _markdown(_decode(content))
    sections = [s for s in sections if s.text.strip()]
    if not sections:
        raise UnsupportedFile(f"no text could be extracted from {filename}")
    title = next((s.heading for s in sections if s.heading), Path(filename).stem)
    return Document(doc_id=doc_id_for(content), source=filename, title=title, sections=sections)


def _decode(content: bytes) -> str:
    for enc in ("utf-8-sig", "utf-16", "cp1252"):
        try:
            return content.decode(enc)
        except UnicodeDecodeError:
            continue
    return content.decode("utf-8", errors="replace")


_HEADING = re.compile(r"^(#{1,4})\s+(.*?)\s*#*\s*$")


def _markdown(text: str) -> list[Section]:
    """Split on headings so every chunk knows which section it came from."""
    sections, heading, buf = [], None, []
    for line in text.splitlines():
        m = _HEADING.match(line)
        if m:
            if buf:
                sections.append(Section("\n".join(buf).strip(), heading))
            heading, buf = m.group(2).strip(), []
        else:
            buf.append(line)
    if buf:
        sections.append(Section("\n".join(buf).strip(), heading))
    return sections


def _pdf(content: bytes) -> list[Section]:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(content))
    return [Section((page.extract_text() or "").strip(), page=i + 1) for i, page in enumerate(reader.pages)]


def _docx(content: bytes) -> list[Section]:
    import docx

    d = docx.Document(io.BytesIO(content))
    sections, heading, buf = [], None, []
    for para in d.paragraphs:
        style = (para.style.name if para.style is not None else "").lower()
        if style.startswith("heading") or style == "title":
            if buf:
                sections.append(Section("\n".join(buf), heading))
            heading, buf = para.text.strip(), []
        elif para.text.strip():
            buf.append(para.text.strip())
    if buf:
        sections.append(Section("\n".join(buf), heading))
    return sections


def _csv(content: bytes) -> list[Section]:
    rows = list(csv.reader(io.StringIO(_decode(content))))
    if not rows:
        return []
    header, body = rows[0], rows[1:]
    # one "sentence" per row, keyed by the header, so rows stay retrievable on their own
    lines = ["; ".join(f"{h}: {v}" for h, v in zip(header, row) if v) for row in body]
    return [Section("\n".join(lines[i : i + 25])) for i in range(0, len(lines), 25)]
